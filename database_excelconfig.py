import sqlite3
import random
import os
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

# Resident and Security details
RESIDENTS = {
    32: ('John Doe', 'B1 L1 Phase 1 Hydrogen'),
    33: ('Jane Smith', 'B1 L2 Phase 1 Helium'),
    34: ('Alice Johnson', 'B1 L3 Phase 1 Lithium'),
    35: ('Michael Brown', 'B1 L4 Phase 1 Beryllium'),
    36: ('Chloe Davis', 'B1 L5 Phase 1 Boron'),
    37: ('Lucas Miller', 'B1 L6 Phase 1 Carbon'),
    38: ('Emma Wilson', 'B1 L7 Phase 1 Nitrogen'),
    39: ('Oliver Moore', 'B1 L8 Phase 1 Oxygen'),
    40: ('Sophia Young', 'B1 L9 Phase 1 Fluorine'),
    41: ('Ethan Johnson', 'B1 L10 Phase 1 Neon'),
    42: ('Mia Williams', 'B1 L11 Phase 1 Hydrogen'),
    43: ('Noah Jones', 'B1 L12 Phase 1 Helium'),
    44: ('Isabella Taylor', 'B1 L13 Phase 1 Lithium'),
    45: ('William White', 'B1 L14 Phase 1 Beryllium'),
    46: ('Ava Thompson', 'B1 L15 Phase 1 Boron'),
    47: ('Matthew Harris', 'B1 L16 Phase 1 Carbon'),
    48: ('Amelia Martin', 'B1 L17 Phase 1 Nitrogen'),
    49: ('James Lee', 'B1 L18 Phase 1 Oxygen'),
    50: ('Charlotte Hall', 'B2 L1 Phase 2 Fluorine'),
    51: ('Alexander Allen', 'B2 L2 Phase 2 Neon'),
    52: ('Harper Young', 'B2 L3 Phase 2 Hydrogen'),
    53: ('Elijah Scott', 'B2 L4 Phase 2 Helium'),
    54: ('Isabelle Edwards', 'B2 L5 Phase 2 Lithium'),
    55: ('Jack Wright', 'B2 L6 Phase 2 Beryllium'),
    56: ('Lily King', 'B2 L7 Phase 2 Boron'),
    57: ('Benjamin Moore', 'B2 L8 Phase 2 Carbon'),
    58: ('Zoe Miller', 'B2 L9 Phase 2 Nitrogen'),
    59: ('Logan Brown', 'B2 L10 Phase 2 Oxygen'),
    60: ('Grace Davis', 'B3 L1 Phase 1 Fluorine'),
    61: ('Ryan Wilson', 'B3 L2 Phase 1 Neon'),
    62: ('Jackson Wong', 'B3 L3 Phase 1 Hydrogen'),
    63: ('Ella Thompson', 'B3 L4 Phase 1 Helium'),
    64: ('Henry Anderson', 'B3 L5 Phase 1 Lithium'),
    65: ('Scarlett Lee', 'B3 L6 Phase 1 Beryllium'),
    66: ('David Martin', 'B3 L7 Phase 1 Boron'),
    67: ('Victoria Clark', 'B3 L8 Phase 1 Carbon'),
    68: ('Daniel Lewis', 'B3 L9 Phase 1 Nitrogen'),
    69: ('Abigail Walker', 'B3 L10 Phase 1 Oxygen'),
    70: ('Mason Young', 'B3 L11 Phase 1 Fluorine'),
    71: ('Hannah Hill', 'B3 L12 Phase 1 Neon'),
    72: ('Jacob Martinez', 'B3 L13 Phase 1 Hydrogen'),
    73: ('Sofia Wright', 'B3 L14 Phase 1 Helium'),
    74: ('William Robinson', 'B3 L15 Phase 1 Lithium'),
    75: ('Grace Mitchell', 'B3 L16 Phase 1 Beryllium'),
    76: ('James Harris', 'B3 L17 Phase 1 Boron')
}

SECURITY = {
    1: 'Juan Dela Cruz',
    2: 'Hannah Santos'
}

# Generate Random Data
def generate_random_data(num_records):
    names = [
        'John Doe', 'Jane Smith', 'Alice Johnson', 'Michael Brown', 'Chloe Davis',
        'Lucas Miller', 'Emma Wilson', 'Oliver Moore', 'Sophia Young', 'Ethan Johnson',
        'Mia Williams', 'Noah Jones', 'Isabella Taylor', 'William White', 'Ava Thompson',
        'Matthew Harris', 'Amelia Martin', 'James Lee', 'Charlotte Hall', 'Alexander Allen',
        'Harper Young', 'Elijah Scott', 'Isabelle Edwards', 'Jack Wright', 'Lily King',
        'Benjamin Moore', 'Zoe Miller', 'Logan Brown', 'Grace Davis', 'Ryan Wilson',
        'Jackson Wong', 'Ella Thompson', 'Henry Anderson', 'Scarlett Lee', 'David Martin',
        'Victoria Clark', 'Daniel Lewis', 'Abigail Walker', 'Mason Young', 'Hannah Hill',
        'Jacob Martinez', 'Sofia Wright', 'William Robinson', 'Grace Mitchell', 'James Harris'
    ]

    purposes = [
        'Will make a visit.', 'Scheduled appointment.', 'Delivery received.', 'Meeting with staff.',
        'Inspection visit.', 'Maintenance check.', 'Guest arrival.', 'Package pickup.',
        'Service repair.', 'Consultation meeting.', 'Training session.', 'Personal visit.',
        'Vendor visit.', 'Document submission.', 'Repair request.', 'Health checkup.',
        'Equipment setup.', 'Special event visit.', 'Emergency visit.', 'Other.'
    ]

    start_date = datetime(2023, 1, 1)
    end_date = datetime(2024, 7, 5)  # Latest date is July 05, 2024

    random_data = []
    for _ in range(num_records):
        visit_name = random.choice(names)
        res_id = random.randint(32, 76)  # Choose randomly from your resident IDs
        log_purpose = random.choice(purposes)
        log_day = start_date + timedelta(days=random.randint(0, (end_date - start_date).days))
        login_time = datetime(log_day.year, log_day.month, log_day.day, random.randint(0, 23), random.randint(0, 59), random.randint(0, 59)).strftime('%H:%M:%S')
        logout_time = datetime(log_day.year, log_day.month, log_day.day, random.randint(0, 23), random.randint(0, 59), random.randint(0, 59)).strftime('%H:%M:%S')
        log_stat = random.choice([0, 1])
        sec_id = random.choice([1, 2])  # Choose randomly from your security IDs

        random_data.append((visit_name, res_id, log_purpose, log_day.strftime('%Y-%m-%d'), login_time, logout_time, log_stat, sec_id))

    return random_data

# Update Visitor Data in SQLite
def update_visitor_data(data):
    conn = sqlite3.connect('visitor_attendance.db')
    cursor = conn.cursor()

    cursor.execute("DELETE FROM visitor_data")

    for data_entry in data:
        cursor.execute("""
            INSERT INTO visitor_data (visit_name, res_id, log_purpose, log_day, login_time, logout_time, log_stat, sec_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, data_entry)

    conn.commit()
    conn.close()

# Update Visitor Data in Excel
def update_visitor_excel(data):
    COL_NAMES = ['VISITOR NAME', 'DATE', 'LOGIN TIME', 'LOGOUT TIME', 'RESIDENT', 'ADDRESS', 'SECURITY', 'PURPOSE']
    
    # Get the home directory of the current user
    home_directory = os.path.expanduser('~')
    # Set the desktop path for the current user
    desktop_path = os.path.join(home_directory, 'Desktop')
    folder_path = os.path.join(desktop_path, 'Visitor_Attendance')
    
    if not os.path.exists(folder_path):
        os.makedirs(folder_path)
    
    # Dictionary to store data by date
    data_by_date = {}
    
    for data_entry in data:
        visit_name, res_id, log_purpose, log_day, login_time, logout_time, log_stat, sec_id = data_entry
        res_name, res_address = RESIDENTS[res_id]
        sec_name = SECURITY[sec_id]
        
        formatted_data = [visit_name, log_day, login_time, logout_time, res_name, res_address, sec_name, log_purpose]
        
        if log_day not in data_by_date:
            data_by_date[log_day] = []
        data_by_date[log_day].append(formatted_data)
    
    # Create separate Excel files for each date
    for log_day, data_list in data_by_date.items():
        file_path = os.path.join(folder_path, f"{log_day}_VAttendance.xlsx")
        
        try:
            if os.path.exists(file_path):
                workbook = load_workbook(file_path)
                sheet = workbook.active
            else:
                workbook = Workbook()
                sheet = workbook.active
                sheet.append(COL_NAMES)  # Add header if the file is newly created

            for formatted_data in data_list:
                sheet.append(formatted_data)

            # Adjust column widths
            for col_num, col_name in enumerate(COL_NAMES, 1):
                column_letter = get_column_letter(col_num)
                max_length = max(len(str(item)) for item in [col_name] + [formatted_data[col_num-1]]) + 2
                sheet.column_dimensions[column_letter].width = max_length

            workbook.save(file_path)
        except Exception as e:
            print(f"Error: {e}")

# Main Function
def main():
    num_records = 100  # Number of records to generate
    data = generate_random_data(num_records)
    update_visitor_data(data)
    update_visitor_excel(data)

if __name__ == "__main__":
    main()
